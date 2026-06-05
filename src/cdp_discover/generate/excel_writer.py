"""excel_writer — write a styled multi-sheet workbook from rows + column specs.

Layout per sheet:

  Row 1: section banner (merged across all columns of the same group)
  Row 2: friendly column header
  Row 3+: data rows
  Frozen panes at A3, autofilter on row 2, alternating row stripes,
  word-wrap for long JSON cells, sensible auto-widths.
"""
from __future__ import annotations

import datetime as _dt
import os
from itertools import cycle
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .columns.aws import Column
from .common import resolve, stringify

# Stable group -> color mapping. Pick palette-friendly tones that read well in
# both Light and Dark Excel themes. New group names auto-pick from EXTRA_COLORS.
_GROUP_COLORS: Dict[str, str] = {
    "Identity":       "1F4E79",
    "Network":        "2E75B6",
    "Storage":        "548235",
    "Security":       "C65911",
    "FreeIPA":        "7030A0",
    "Tags":           "BF8F00",
    "Environment":    "1F4E79",
    "Datalake":       "385723",
    "Cluster":        "385723",
    "Instance Group": "806000",
    "Instance":       "00536E",
    "Image":          "5B0F00",
    "Owner":          "1F4E79",
    "DB Server":      "385723",
    "SSL":            "7030A0",
    "Workspace":      "1F4E79",
    "Capacity":       "385723",
    "Creator":        "5B0F00",
    "Support":        "C65911",
    "Service":        "1F4E79",
    "Event":          "7030A0",
    "Registry":       "7030A0",
    "App":            "1F4E79",
    "Data Viz":       "385723",
    "Hue":            "806000",
    "DBC":            "385723",
    "VW":             "2E75B6",
    "Backups":        "548235",
    "Access":         "C65911",
    "Upgrade":        "C65911",
    "Virtual Cluster":"385723",
}
_EXTRA_COLORS = cycle(["404040", "595959", "7B7B7B", "0070C0", "203864"])
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_LABEL_FONT = Font(color="FFFFFF", bold=True, size=10)
_LABEL_FILL = PatternFill("solid", fgColor="305496")
_STRIPE_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(border_style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top", horizontal="left")
_CENTER_BOLD = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _color_for(group: str) -> str:
    if group not in _GROUP_COLORS:
        _GROUP_COLORS[group] = next(_EXTRA_COLORS)
    return _GROUP_COLORS[group]


def _segment_groups(columns: Sequence[Column]) -> List[Tuple[str, int, int]]:
    """Return list of (group, start_col_index_1based, end_col_index_1based)."""
    out: List[Tuple[str, int, int]] = []
    if not columns:
        return out
    start = 1
    cur = columns[0].group
    for i, col in enumerate(columns, start=1):
        if col.group != cur:
            out.append((cur, start, i - 1))
            start = i
            cur = col.group
    out.append((cur, start, len(columns)))
    return out


def _autosize(ws: Worksheet, columns: Sequence[Column], rows: Sequence[Sequence]) -> None:
    """Set sensible column widths based on header + first ~20 rows of data."""
    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        max_len = max(12, len(col.label) + 2)
        for r in rows[:20]:
            v = r[idx - 1]
            if v is None:
                continue
            text = str(v)
            line = max(len(s) for s in text.splitlines()) if text else 0
            if line + 2 > max_len:
                max_len = line + 2
        ws.column_dimensions[letter].width = min(max(14, max_len), 60)


def write_sheet(
    wb: Workbook,
    sheet_name: str,
    columns: Sequence[Column],
    records: Sequence[dict],
) -> Worksheet:
    """Render a single sheet, even if ``records`` is empty (writes headers only)."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    if not columns:
        ws.cell(row=1, column=1, value=f"(no columns defined for {sheet_name})")
        return ws

    for group, c_start, c_end in _segment_groups(columns):
        cell = ws.cell(row=1, column=c_start, value=group.upper())
        cell.fill = PatternFill("solid", fgColor=_color_for(group))
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_BOLD
        cell.border = _BORDER
        if c_end > c_start:
            ws.merge_cells(start_row=1, end_row=1, start_column=c_start, end_column=c_end)

    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=idx, value=col.label)
        cell.font = _LABEL_FONT
        cell.fill = _LABEL_FILL
        cell.alignment = _CENTER_BOLD
        cell.border = _BORDER
    ws.row_dimensions[2].height = 30

    rendered_rows: List[List] = []
    for record in records:
        rendered_rows.append([
            stringify(resolve(col.accessor, record)) for col in columns
        ])

    for r_idx, row in enumerate(rendered_rows, start=3):
        is_stripe = (r_idx % 2 == 1)
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = _WRAP
            cell.border = _BORDER
            if is_stripe:
                cell.fill = _STRIPE_FILL
            if isinstance(value, _dt.datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"

    last_col_letter = get_column_letter(len(columns))
    last_row = max(2, 2 + len(rendered_rows))
    ws.auto_filter.ref = f"A2:{last_col_letter}{last_row}"
    ws.freeze_panes = "A3"
    _autosize(ws, columns, rendered_rows)

    if not rendered_rows:
        info = ws.cell(row=3, column=1, value=f"No data found for {sheet_name}.")
        info.font = Font(italic=True, color="777777")
    return ws


def _summary_sheet(wb: Workbook, sheets_info: List[Tuple[str, int, int]]) -> Worksheet:
    """First-tab summary: rows per sheet + columns per sheet."""
    ws = wb.create_sheet("SUMMARY", 0)
    ws.cell(row=1, column=1, value="UPGRADE PLANNING REPORT").font = Font(size=14, bold=True, color="1F4E79")
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=4)
    ws.cell(row=2, column=1, value="Generated").font = Font(bold=True)
    ws.cell(row=2, column=2, value=_dt.datetime.now()).number_format = "yyyy-mm-dd hh:mm:ss"

    headers = ["Sheet", "Rows", "Columns", "Notes"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font = _LABEL_FONT
        cell.fill = _LABEL_FILL
        cell.alignment = _CENTER_BOLD
        cell.border = _BORDER
    for r, (name, n_rows, n_cols) in enumerate(sheets_info, start=5):
        ws.cell(row=r, column=1, value=name).border = _BORDER
        ws.cell(row=r, column=2, value=n_rows).border = _BORDER
        ws.cell(row=r, column=3, value=n_cols).border = _BORDER
        ws.cell(row=r, column=4, value="" if n_rows else "no data found").border = _BORDER

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 38
    ws.freeze_panes = "A5"
    return ws


def write_workbook(
    output_path: str,
    sheets: Iterable[Tuple[str, Sequence[Column], Sequence[dict]]],
) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    info: List[Tuple[str, int, int]] = []
    for name, columns, records in sheets:
        write_sheet(wb, name, columns, records)
        info.append((name, len(records), len(columns)))
    _summary_sheet(wb, info)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)
    wb.save(output_path)
    return output_path
